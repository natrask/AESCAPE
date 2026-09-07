"""Build slides/agentic_ai_intro.pptx -- an editable port of the beamer deck."""
import os
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from deck_lib import (Deck, ML, MR, CW, SW, SH, BODY_TOP, BODY_BOT, TITLE_TOP,
                      INK, GRAY, LGRAY, ACCENT, WHITE, EDGE,
                      FILL_BLUE, FILL_ORANGE, FILL_GREEN, FILL_PURPLE, FILL_GRAY,
                      FILL_CODE, FILL_NOTE, FONT, MONO, add_runs, set_space)

EQ = os.path.join(os.path.dirname(os.path.abspath(__file__)), "eq", "indicator.png")
d = Deck()

# ===========================================================  1. TITLE
s = d.slide(footer=False)
tb = s.shapes.add_textbox(Inches(ML), Inches(2.15), Inches(CW), Inches(1.9))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; add_runs(p, "Foundations of Agentic AI", size=42, bold=True)
p2 = tf.add_paragraph(); add_runs(p2, "for Science & Engineering", size=42, bold=True)
sub = s.shapes.add_textbox(Inches(ML), Inches(4.15), Inches(CW), Inches(1.6))
tf = sub.text_frame; tf.word_wrap = True
for i, (txt, sz, col) in enumerate([
        ("AESCAPE 2026  ·  SC03: Introduction to Agentic Workflows", 17, ACCENT),
        ("", 8, GRAY),
        ("Nathaniel Trask", 16, INK),
        ("University of Pennsylvania", 13, GRAY),
        ("September 8, 2026", 13, GRAY)]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    add_runs(p, txt, size=sz, color=col); set_space(p, after=3)
ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(ML), Inches(4.02),
                            Inches(ML + 3.4), Inches(4.02))
ln.line.color.rgb = ACCENT; ln.line.width = Pt(2.5)

# ===========================================================  2. TODAY
s = d.slide("Today", notes="Frame the whole talk around the one question: who owns control flow.")
d.body(s, [
    "1.  **Foundations.**  What an agent is. Tools, schemas, the loop. Build ReAct from scratch in ~30 lines.",
    "",
    "2.  **The framework zoo.**  ReAct, graphs (LangGraph), conversational multi-agent (AutoGen), MPC-style planning — and MCP.  *Syntax, a toy, and an exemplar for each.*",
    "",
    "3.  **Scientific computing as a tool surface.**  CAD → mesh → solve → post, re-cut so an LLM can drive it.",
    "",
    "4.  **Reliability.**  Verifier agents, manufactured solutions, oracles.",
], size=17, spacing=4)
d.note(s, "**The one question this talk is organized around:**  for any given problem, "
          "!!which agentic pattern should you reach for, and why?!!  Not “how do I use library X.”",
       y=5.45, size=15)

# ===========================================================  3. WHAT IS AN AGENT
s = d.slide("What is an LLM agent?")
d.label(s, ML, 1.22, CW, "Agent  =  LLM  +  tools  +  a loop.", size=24, color=INK,
        align=PP_ALIGN.CENTER)
d.body(s, [
    "- A bare LLM takes text in, emits text out. One shot.",
    "- An !!agent!! is an LLM running in a loop with access to !!tools!!: ordinary Python functions it can ask us to run on its behalf.",
    "- Each turn, the model reads the running conversation and decides to either",
    "  - emit a final text answer, or",
    "  - call a tool with structured arguments.",
    "- When it calls a tool, *we* run the function and feed the result back as the next turn.",
], y=2.05, size=16.5)
d.note(s, "The model never executes anything. It emits a *request*; you stay in control of what "
          "actually runs.", y=5.5)

# ===========================================================  4. THE LOOP (diagram)
s = d.slide("The loop, as a picture")
llm = d.box(s, 2.25, 2.55, 3.3, 1.35, "LLM\nreads messages,\nemits text or tool call", FILL_BLUE, size=13)
tls = d.box(s, 7.75, 2.55, 3.3, 1.35, "Tools\nPython functions\nwe control", FILL_GREEN, size=13)
a1 = d.arrow(s, 5.55, 2.85, 7.75, 2.85)
a2 = d.arrow(s, 7.75, 3.60, 5.55, 3.60)
d.label(s, 5.35, 2.18, 2.7, "function_call\n(name, args)", size=10)
d.label(s, 5.35, 3.72, 2.7, "result (dict)\nappended as observation", size=10)
d.label(s, ML, 4.55, CW, "when the model stops calling tools, it emits a final text answer and the loop exits",
        size=11, italic=True)
d.label(s, ML, 5.35, CW, "This is !!ReAct!! (Reason + Act), Yao et al. 2022 — the simplest agent there is.",
        size=16, color=INK)

# ===========================================================  5. WHAT A TOOL IS
s = d.slide("What a tool is")
d.body(s, ["Two pieces, both written by you: a Python function, and a JSON-schema description "
           "of its signature that you hand to the model."], size=15)
d.label(s, ML, 1.68, 5.9, "Python function  (you call it)", size=12, color=GRAY, align=PP_ALIGN.LEFT)
d.code(s, '''
def calculator(op, a, b):
    ops = {"add": a + b, "sub": a - b,
           "mul": a * b, "div": a / b}
    return {"result": ops[op]}
''', x=ML, y=1.98, w=5.9, size=11.5)
d.label(s, ML + 6.2, 1.68, 5.9, "JSON schema  (the model reads it)", size=12, color=GRAY, align=PP_ALIGN.LEFT)
d.code(s, '''
{"name": "calculator",
 "description": "One arithmetic op.",
 "parameters": {
   "type": "object",
   "properties": {
     "op": {"type": "string",
            "enum": ["add","sub","mul","div"]},
     "a": {"type": "number"},
     "b": {"type": "number"}},
   "required": ["op","a","b"]}}
''', x=ML + 6.2, y=1.98, w=5.9, size=11.5)
d.body(s, [
    "- `name`, `description` — decide whether the model picks *this* tool over the others. This is prompt engineering; write it carefully.",
    "- `properties` — the typed argument list. `enum` restricts a string to a fixed set. Constrain aggressively: every value you forbid in the schema is a failure mode you never have to debug.",
], y=4.35, size=14)

# ===========================================================  6. ROUND TRIP
s = d.slide("The round trip")
d.body(s, ["Each turn, the model emits either plain text (final answer) or a structured tool call."], size=15)
d.body(s, ["**1.**  The reply contains a `function_call` part:"], y=1.72, size=14.5)
d.code(s, '''
response.candidates[0].content.parts[0].function_call
    .name = "calculator"
    .args = {"op": "mul", "a": 13, "b": 47}
''', y=2.10, size=12)
d.body(s, ["**2.**  We dispatch by name and run the Python function:"], y=3.28, size=14.5)
d.code(s, '''
result = tool_fns[name](**args)          # -> {"result": 611}
''', y=3.66, size=12)
d.body(s, ["**3.**  We send the result back as a `function_response` on the next turn:"], y=4.42, size=14.5)
d.code(s, '''
Part.from_function_response(name="calculator", response={"result": 611})
''', y=4.80, size=12)
d.body(s, ["The model then sees its own call, the result, and decides whether to call another "
           "tool or stop."], y=5.60, size=14.5)

# ===========================================================  7. PSEUDOCODE
s = d.slide("The whole thing, in pseudocode")
d.code(s, '''
def run_agent(system_prompt, user_prompt, tool_schemas, tool_fns, max_steps=12):
    messages = [user_prompt]
    for step in range(max_steps):
        response = llm(system_prompt, messages, tools=tool_schemas)
        messages.append(response)

        if response.is_tool_call:
            for call in response.tool_calls:
                try:
                    result = tool_fns[call.name](**call.args)
                except Exception as e:
                    result = {"error": str(e)}      # errors are observations
                messages.append(result)
        else:
            return response.text                    # final answer
''', y=1.30, size=13)
d.label(s, ML, 5.55, CW,
        "That is a complete agent. Everything else in this talk is a variation on "
        "!!who decides what happens next!!.", size=17, color=INK, align=PP_ALIGN.LEFT)

# ===========================================================  8. TWO SLOTS
s = d.slide("Two slots: system_prompt vs user_prompt")
d.body(s, [
    "- !!system_prompt!! — the *rules of the game*. Identity, methodology, what the tools are for, budgets, how to stop. Sent on every turn as a persistent header; never drowned out as observations pile up.",
    "- !!user_prompt!! — the *specific ask* that triggers this run. Becomes the first message; tool calls and observations append after it.",
], size=15.5)
d.code(s, '''
B1_SYSTEM = """You optimize a black-box function (angle in degrees, in (0,90))
returning a range in meters. Budget: at most 12 evaluate() calls. Use history if
helpful, then call propose_answer with your guess."""

B1_USER = "Find the optimal angle. Budget: 12 evaluations."
''', y=3.30, size=12)
d.note(s, "Rule of thumb: durable rules about *how to behave* go in the system prompt; the "
          "specific *thing to do this run* goes in the user prompt.", y=5.35)

# ===========================================================  9. ERRORS
s = d.slide("Errors are observations, not crashes")
d.body(s, ["Tool calls **will** fail. The model invents an argument, asks for a mesh ID that "
           "does not exist, passes a string where you wanted a float."], size=15.5)
d.code(s, '''
try:
    result = tool_fns[name](**args)
except Exception as e:
    result = {"error": f"{type(e).__name__}: {e}"}
''', y=2.20, size=12.5)
d.body(s, ["The model reads the error on the next turn and, with a reasonable prompt, retries "
           "with correct arguments."], y=3.75, size=15.5)
d.note(s, "**You will see this happen live today.**  An agent recovering from its own bad tool "
          "call is the system working, not the system failing. The failure mode to actually fear "
          "is on the slide “agents are confidently wrong.”", y=4.85, size=14)

# ===========================================================  10. GUARDRAILS
s = d.slide("Guardrails belong in the tool, not the prompt")
d.body(s, ["A prompt is a request. A tool is an enforcement point."], size=16)
d.table(s, [
    ["Hope  (prompt)", "Guarantee  (tool)"],
    ["“don’t refine past 20k DOFs”", "`if new_dofs > MAX_DOFS: return {\"error\": ...}`"],
    ["“use at most 12 evaluations”", "counter in the tool; refuse call 13"],
    ["“don’t touch the boundary”", "the tool has no boundary argument"],
    ["“stop when converged”", "`stop()` is the only terminal tool"],
], y=1.95, col_w=[1, 1.5], size=13.5, row_h=0.42)
d.body(s, [
    "Anything you would be unwilling to let a nondeterministic process do, !!make impossible in "
    "the tool signature!! — do not ask nicely in English.",
    "",
    "This is the single highest-leverage habit for making agentic workflows safe enough to run "
    "unattended.",
], y=4.35, size=15)

# ===========================================================  11. SETUP
s = d.slide("Setup: getting a key")
d.body(s, [
    "1.  **Free, no credit card.**  Google AI Studio: aistudio.google.com/apikey. Sign in with any Google account. Use a project where billing has *never* been enabled.",
    "",
    "2.  **Bring your own.**  Anthropic (console.anthropic.com) or OpenAI (platform.openai.com/api-keys) both work.",
    "",
    "3.  In Colab: **Secrets** panel (key icon, left sidebar), add `GEMINI_API_KEY`. Locally: `export GEMINI_API_KEY=...`",
    "",
    "4.  Run the Part 0 cells to confirm the client connects.",
], size=15.5, spacing=3)
d.note(s, "**Never paste a key into a notebook cell.**  It ends up in the `.ipynb` JSON, in your "
          "shell history, and in git. Use the secrets panel or an environment variable.", y=5.35)

# ===========================================================  12. SECTION: ZOO
def divider(title, subtitle=None, notes=None):
    sl = d.slide(footer=False, notes=notes)
    tb = sl.shapes.add_textbox(Inches(ML), Inches(2.9), Inches(CW), Inches(1.4))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    add_runs(p, title, size=38, bold=True)
    if subtitle:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        set_space(p2, before=10)
        add_runs(p2, subtitle, size=17, color=GRAY)
    ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(5.9), Inches(2.62),
                                 Inches(7.43), Inches(2.62))
    ln.line.color.rgb = ACCENT; ln.line.width = Pt(2.5)
    return sl

divider("The framework zoo", "ReAct  ·  LangGraph  ·  AutoGen  ·  MPC  ·  MCP")

# ===========================================================  13. WHY NOT REACT
s = d.slide("Why not just ReAct?")
d.body(s, [
    "ReAct is **greedy**. At each step the model sees the history and picks one action. That is exactly right when:",
    "- the tool surface is small,",
    "- mistakes are cheap and recoverable, and",
    "- you cannot specify the procedure in advance.",
], size=15.5)
d.body(s, [
    "It starts to hurt when:",
    "- you need a !!guarantee!! that step B always follows step A (audit, compliance, “always verify before reporting”),",
    "- genuinely !!different expertise!! needs different system prompts and different tools,",
    "- a wrong action is !!expensive!! and you own a cheap forward model that could have caught it, or",
    "- your tools live in !!someone else’s process!! — a solver, a CAD kernel, a database.",
], y=3.45, size=15.5)
d.label(s, ML, 6.20, CW, "Each of those pressures produced a different pattern. That is the zoo.",
        size=15, color=INK, align=PP_ALIGN.LEFT)

# ===========================================================  14. CONTROL FLOW SPECTRUM
s = d.slide("The only question that matters")
d.label(s, ML, 1.35, CW, "Choosing a pattern is choosing !!who owns control flow!!.",
        size=22, color=INK, align=PP_ALIGN.CENTER)
specs = [("ReAct", "the *model*\ndecides", FILL_BLUE),
         ("LangGraph", "*you* decide,\nin code", FILL_ORANGE),
         ("AutoGen", "the *conversation*\ndecides", FILL_PURPLE),
         ("MPC", "a *simulator*\ndecides", FILL_GREEN)]
bx, bw, gap, by, bh = ML, 2.10, 0.28, 2.45, 1.35
for i, (nm, sub, fill) in enumerate(specs):
    d.box(s, bx + i * (bw + gap), by, bw, bh, f"{nm}\n{sub}", fill, size=12.5)
right_edge = bx + 4 * bw + 3 * gap
div = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(right_edge + 0.16), Inches(2.25),
                             Inches(right_edge + 0.16), Inches(4.05))
div.line.color.rgb = LGRAY; div.line.width = Pt(1.0)
div.line.dash_style = 4
d.box(s, right_edge + 0.42, by, bw, bh, "MCP\n*nobody* —\nit’s plumbing", FILL_GRAY, size=12.5)
d.arrow(s, ML - 0.05, 4.22, right_edge - 0.10, 4.22, tail=True, color=LGRAY)
d.label(s, ML - 0.10, 4.35, 2.0, "more autonomy", size=11, align=PP_ALIGN.LEFT)
d.label(s, right_edge - 2.0, 4.35, 2.0, "more control", size=11, align=PP_ALIGN.RIGHT)
d.label(s, right_edge + 0.42, 4.30, bw, "a *transport*,\nnot an architecture", size=10)
d.body(s, ["MCP sits apart on purpose: it is a *transport*, not an architecture. You can use it "
           "with any of the other four. (More on the unfortunate MPC/MCP name collision shortly.)"],
       y=5.30, size=14.5)

# ===========================================================  15. ZOO TABLE
s = d.slide("The zoo, on one slide")
d.table(s, [
    ["", "Control flow", "Lookahead", "Roles", "Reach for it when"],
    ["**ReAct**", "model, step by step", "none", "one", "open-ended, errors cheap"],
    ["**LangGraph**", "you, as a graph", "none", "many, fixed", "you need guarantees"],
    ["**AutoGen**", "emergent, via chat", "none", "many, fluid", "expertise must collide"],
    ["**MPC**", "simulator ranks plans", "*N* steps", "one + model", "wrong actions are costly"],
    ["**MCP**", "*orthogonal* — how tools reach the agent, not how it decides", "", "", ""],
], y=1.35, col_w=[1.0, 1.85, 0.85, 1.0, 1.85], size=13, row_h=0.45)
d.body(s, [
    "For each of the four architectures we will look at:",
    "- the **syntax** — what the code actually looks like,",
    "- a **toy** — the *same* trivial task, so the comparison is fair,",
    "- an **exemplar** — a problem where that pattern genuinely wins.",
], y=4.35, size=15.5)

# ===========================================================  16. COMMON TOY
s = d.slide("The common toy task")
d.body(s, ["To compare syntax fairly, every framework gets the *same* trivial job:"], size=16)
d.note(s, "**Toy task.**  Given a target number, use a `calculator` tool to reach it, then report "
          "the result. No framework should need more than ~20 lines.", y=1.95, x=1.6, w=CW - 2.0, size=15)
d.body(s, [
    "The toy is deliberately **too easy to justify any of them**. That is the point: when the task "
    "is trivial, all four look like unnecessary ceremony, and you can see the raw syntax without "
    "the problem getting in the way.",
    "",
    "Then each **exemplar** shows the task where the ceremony pays for itself.",
], y=3.20, size=16)
d.label(s, ML, 5.85, CW, "Notebook: 01_agentic_patterns.ipynb, §1–§5", size=12, align=PP_ALIGN.CENTER)

# ===========================================================  17-19. REACT
divider("Pattern 1 — ReAct", "the model owns control flow")

s = d.slide("ReAct — syntax")
d.body(s, ["No framework. A `for` loop and a dispatch table."], size=15)
d.code(s, '''
contents = [user_prompt]
for step in range(max_steps):
    resp = client.models.generate_content(
        model=MODEL, contents=contents,
        config=GenerateContentConfig(
            system_instruction=system_prompt,
            tools=[Tool(function_declarations=tool_schemas)]))
    contents.append(resp.candidates[0].content)

    calls = [p.function_call for p in resp.candidates[0].content.parts
             if p.function_call]
    if not calls:
        return resp.text                       # final answer

    for fc in calls:
        result = tool_fns[fc.name](**dict(fc.args))
        contents.append(Part.from_function_response(
            name=fc.name, response=result))
''', y=1.65, size=12)
d.label(s, ML, 6.15, CW, "Dependencies: one SDK.  State: the `contents` list.  That’s it.",
        size=13, color=INK, align=PP_ALIGN.LEFT)

s = d.slide("ReAct — exemplar: adaptive mesh refinement")
d.body(s, [
    "**Why ReAct is the right call here:**",
    "- The number of refinement cycles is !!not knowable in advance!! — it depends on observed error.",
    "- Every action is !!cheap and recoverable!!: a bad refinement just makes a mesh you discard.",
    "- The model must !!react to observations!! it could not have predicted (the error estimate after each solve).",
    "- The tool surface is !!small!! — six tools.",
    "",
    "Greedy, one-step-at-a-time decision making is not a compromise here. It is the correct algorithm.",
], w=7.1, size=14.5)
bx, bw, bh = 8.55, 2.9, 0.55
n1 = d.box(s, bx, 1.55, bw, bh, "solve_poisson", FILL_BLUE, size=11.5, bold_first=False)
n2 = d.box(s, bx, 2.45, bw, bh, "local_indicators", FILL_BLUE, size=11.5, bold_first=False)
n3 = d.box(s, bx, 3.35, bw, bh, "refine_by_threshold", FILL_BLUE, size=11.5, bold_first=False)
n4 = d.box(s, bx, 4.55, bw, bh, "stop", FILL_GRAY, size=11.5, bold_first=False)
d.connect(s, n1, n2, 2, 0)
d.connect(s, n2, n3, 2, 0)
d.connect(s, n3, n4, 2, 0)
d.arrow(s, bx - 0.30, 3.62, bx - 0.30, 1.82, kind=MSO_CONNECTOR.ELBOW)
d.arrow(s, bx, 3.62, bx - 0.30, 3.62, head=False)
d.arrow(s, bx - 0.30, 1.82, bx, 1.82)
d.label(s, bx - 1.55, 2.55, 1.2, "loop", size=10, align=PP_ALIGN.RIGHT)
d.label(s, bx + bw + 0.05, 4.02, 1.3, "tol met", size=10, align=PP_ALIGN.LEFT)
d.label(s, bx - 0.2, 5.35, bw + 0.4,
        "The model chooses the arrow to take at every step. Nothing in the code forces this order.",
        size=10.5, italic=True)

# ===========================================================  20-23. LANGGRAPH
divider("Pattern 2 — Graphs (LangGraph)", "you own control flow")

s = d.slide("LangGraph — syntax")
d.body(s, ["You declare a typed state, nodes that transform it, and edges. The LLM fills in "
           "nodes; *the graph* decides what runs next."], size=14.5)
d.code(s, '''
from langgraph.graph import StateGraph, END
from typing import TypedDict

class S(TypedDict):
    mesh_id: int
    error: float
    attempts: int

def solve(s):  return {"error": run_solver(s["mesh_id"]), ...}
def refine(s): return {"mesh_id": do_refine(s["mesh_id"]),
                       "attempts": s["attempts"] + 1}

def route(s):                                  # <- control flow, in Python
    if s["error"] < TOL:      return "verify"
    if s["attempts"] >= 6:    return "give_up"
    return "refine"

g = StateGraph(S)
g.add_node("solve", solve); g.add_node("refine", refine); g.add_node("verify", verify)
g.set_entry_point("solve")
g.add_conditional_edges("solve", route,
    {"refine": "refine", "verify": "verify", "give_up": END})
g.add_edge("refine", "solve")
app = g.compile()
''', y=1.90, size=11.5)

s = d.slide("LangGraph — what you actually bought")
d.body(s, [
    "**Guarantees, not suggestions:**",
    "- `verify` !!always!! runs before a result is reported. The model cannot decide to skip it.",
    "- The retry count is !!bounded in code!!. No prompt can talk it into a seventh attempt.",
    "- State is !!typed and inspectable!! — you can log, checkpoint, and resume mid-run.",
    "- The control flow is a !!diagram you can show a reviewer!!, which matters when the answer feeds a design decision.",
], w=6.55, size=14.5)
# A clean vertical cycle:  solve -> verify -> (FAIL) -> refine -> solve
CX, RX, BW = 10.15, 7.45, 1.85
g1 = d.box(s, CX, 1.45, BW, 0.55, "solve",  FILL_BLUE,   size=12)
g3 = d.box(s, CX, 2.70, BW, 0.55, "verify", FILL_GREEN,  size=12)
g4 = d.box(s, CX, 4.15, BW, 0.55, "END",    FILL_GRAY,   size=12)
g2 = d.box(s, RX, 2.70, BW, 0.55, "refine", FILL_ORANGE, size=12)
d.connect(s, g1, g3, 2, 0)                          # solve -> verify (always)
d.connect(s, g3, g4, 2, 0)                          # verify -> END
d.connect(s, g3, g2, 1, 3, color=ACCENT)            # verify -> refine (FAIL)
d.arrow(s, RX + BW / 2, 2.70, RX + BW / 2, 1.72, head=False)   # refine up...
d.arrow(s, RX + BW / 2, 1.72, CX, 1.72)                        # ...back into solve
d.label(s, RX + BW, 2.42, 0.70, "FAIL", size=10, color=ACCENT)
d.label(s, RX - 1.05, 1.98, 1.0, "retry", size=10, align=PP_ALIGN.RIGHT)
d.label(s, RX - 0.15, 5.05, 5.0,
        "Compare with the ReAct diagram: same boxes, but here the *arrows are code*. "
        "The model never sees them.", size=10.5, italic=True)

s = d.slide("LangGraph — exemplar: solve → verify → repair")
d.body(s, ["**The problem:** you want an FEM result you can put in a report."], size=16)
d.table(s, [
    ["Requirement", "Who enforces it"],
    ["Every reported result is verified", "the graph, not the prompt"],
    ["At most 6 refinement attempts", "the router, in Python"],
    ["A failed verification triggers repair, not a retry", "a conditional edge"],
    ["The whole run is reproducible from a checkpoint", "typed state"],
], y=2.00, col_w=[1.5, 1.0], size=14, row_h=0.42)
d.body(s, [
    "None of these are things you want a language model deciding on a given Tuesday. When “it "
    "usually does the right thing” is not good enough, you stop asking and start wiring.",
], y=4.30, size=15.5)
d.label(s, ML, 5.45, CW, "Reach for a graph when the control flow is a !!requirement!!, not a choice.",
        size=18, color=INK, align=PP_ALIGN.CENTER)

# ===========================================================  24-26. AUTOGEN
divider("Pattern 3 — Conversational multi-agent (AutoGen)",
        "the conversation owns control flow")

s = d.slide("AutoGen — syntax")
d.body(s, ["You define participants with distinct system prompts and let them talk. Control flow "
           "emerges from who speaks next."], size=14.5)
d.code(s, '''
from autogen_agentchat.agents import AssistantAgent
from autogen_agentchat.teams import RoundRobinGroupChat
from autogen_agentchat.conditions import TextMentionTermination

modeler = AssistantAgent("modeler", model_client=client,
    system_message="You propose discretizations. Be concrete and specific.")

analyst = AssistantAgent("analyst", model_client=client,
    system_message="""You are a numerical analyst. Challenge the modeler's
    proposal on stability, conditioning, and convergence rate. Do not be
    agreeable; your value is in the objection.""")

critic = AssistantAgent("critic", model_client=client,
    system_message="Judge the exchange. Say APPROVE only when convinced.")

team = RoundRobinGroupChat([modeler, analyst, critic],
    termination_condition=TextMentionTermination("APPROVE"))

await team.run(task="Choose a discretization for advection-dominated flow.")
''', y=1.90, size=11.5)

s = d.slide("AutoGen — exemplar: the design review")
d.body(s, [
    "**The problem:** choosing a discretization for an advection-dominated problem. There is no "
    "single correct answer, and the *reasoning* is the deliverable.",
    "",
    "- A single ReAct agent asked “pick a discretization” produces a confident paragraph with no adversarial pressure on it.",
    "- Three agents with !!genuinely different objectives!! — propose, object, judge — produce a transcript in which the weaknesses are stated out loud.",
    "- The artifact is not the answer. It is !!the argument!!, which a human can audit far faster than they can re-derive.",
], size=15)
d.label(s, ML, 4.60, CW, "Reach for conversation when !!disagreement is the product!!.",
        size=18, color=INK, align=PP_ALIGN.CENTER)
d.note(s, "Caveat worth stating plainly: this pattern is the easiest of the four to fool yourself "
          "with. Three LLMs agreeing is not three experts agreeing — they share a prior. Distinct "
          "*tools* per role helps far more than distinct adjectives in the system prompt.",
       y=5.35, size=13)

# ===========================================================  27-31. MPC
divider("Pattern 4 — MPC-style planning", "a simulator owns control flow")

s = d.slide("MPC — the idea, borrowed from control")
d.body(s, ["Model Predictive Control, as you already know it:"], size=15)
mb = [("state sₜ", FILL_GRAY, ""), ("plan", FILL_BLUE, "N candidates"),
      ("simulate", FILL_GREEN, "forward model"), ("act", FILL_ORANGE, "first step only")]
bw, gap, bx0, by = 2.55, 0.55, 0.95, 1.95
shapes_ = []
for i, (nm, fill, sub) in enumerate(mb):
    txt = f"{nm}\n{sub}" if sub else nm
    shapes_.append(d.box(s, bx0 + i * (bw + gap), by, bw, 0.95, txt, fill, size=13))
for i in range(3):
    d.connect(s, shapes_[i], shapes_[i + 1], 3, 1)
_ac, _sc, _yb, _yl = bx0 + 3 * (bw + gap) + bw / 2, bx0 + bw / 2, 2.90, 3.48
d.arrow(s, _ac, _yb, _ac, _yl, head=False)      # down from "act"
d.arrow(s, _ac, _yl, _sc, _yl, head=False)      # left along the bottom
d.arrow(s, _sc, _yl, _sc, _yb)                  # up into "state"
d.label(s, ML, 3.54, CW, "re-plan from the new state", size=11)
d.body(s, [
    "The agentic version swaps one box: !!the LLM proposes the candidate plans!!, and the "
    "simulator — not the LLM — ranks them.",
    "",
    "- The LLM contributes what it is good at: *proposing plausible candidates* from context.",
    "- The simulator contributes what it is good at: *being right*.",
    "- You commit only to the first action, then re-plan with real feedback.",
], y=4.15, size=15)

s = d.slide("MPC — syntax")
d.body(s, ["No library needed. The pattern is the loop."], size=15)
d.code(s, '''
state = initial_state
for t in range(horizon):
    # 1. LLM proposes N candidate action sequences (cheap, creative)
    plans = llm_propose(state, history, n_candidates=5, lookahead=3)

    # 2. The *simulator* scores each one (trusted, not creative)
    scores = [world_model.rollout(state, plan) for plan in plans]

    # 3. Commit to the first action of the best plan only
    best = plans[argmax(scores)]
    state = env.step(best[0])            # <- the only real action taken

    # 4. Re-plan from the new state (closes the loop on reality)
''', y=1.80, size=12.5)
d.body(s, ["Note what the LLM never does: !!it never decides which plan is best!!. It generates "
           "hypotheses; the forward model adjudicates. That division is the entire value of the "
           "pattern."], y=4.95, size=15.5)

s = d.slide("MPC — exemplar: black-box optimization",
            notes="Run the head-to-head cell live if time allows. Expect ReAct to vary run to run; "
                  "MPC is much more stable because the surrogate does the deciding.")
d.body(s, ["**The problem:** find the launch angle maximizing range for a projectile under "
           "quadratic drag. The agent sees only `evaluate(angle)`."], size=15)
d.code(s, '''
x'' = -c_d |v| x'          y'' = -g - c_d |v| y'
v0 = 50 m/s      c_d = 0.01 1/m      g = 9.81 m/s^2
''', y=2.00, size=13.5)
d.body(s, [
    "- Without drag the optimum is 45°; with drag it shifts lower. The model’s prior is *almost* right, which is the interesting case.",
    "- Every `evaluate` call is a real ODE integration — the thing we are budgeting.",
    "- We have a !!cheap, trusted forward model!!: a surrogate fitted to what we have seen.",
], y=3.05, size=15)
d.label(s, ML, 4.85, CW,
        "In the notebook we run !!ReAct vs. MPC head to head on the same budget!! and count "
        "evaluations to a fixed tolerance.", size=16, color=INK, align=PP_ALIGN.CENTER)
d.note(s, "This is the honest way to answer “when would I reach for one over the other” — "
          "measure it, don’t assert it.", y=5.55, size=13)

s = d.slide("MPC — when it wins, and when it doesn’t")
d.body(s, [
    "**Reach for MPC when:**",
    "- A wrong action is !!expensive or irreversible!! — burning budget, committing a mesh, running an experiment.",
    "- You have a !!cheap forward model you trust!! more than the LLM.",
    "- Actions compose, so looking *N* steps ahead genuinely differs from looking one step ahead.",
], w=6.0, size=14)
d.body(s, [
    "**Don’t bother when:**",
    "- Simulating a plan costs about what executing it costs — then just execute and observe.",
    "- You have no forward model. Then MPC degenerates into the LLM grading its own homework, which is worse than ReAct because it is *slower* and equally wrong.",
    "- The problem is convex and small. Use `scipy.optimize` and go home.",
], x=6.95, w=5.75, size=14)
d.note(s, "**The uncomfortable baseline.**  For the projectile, `minimize_scalar` finds the "
          "optimum in ~19 evaluations with no LLM at all. We run it in the notebook. If an agent "
          "cannot beat scipy, that is worth knowing — and worth saying out loud.", y=5.35, size=13.5)

# ===========================================================  32-37. MCP
divider("Pattern 5 — MCP", "not an architecture at all")

s = d.slide("MPC is not MCP")
d.label(s, ML, 1.30, CW, "Two three-letter acronyms, one letter apart, both in this talk. Sorry.",
        size=18, color=INK, align=PP_ALIGN.CENTER)
d.table(s, [
    ["MPC — Model Predictive Control", "MCP — Model Context Protocol"],
    ["A *control strategy*.", "A *wire protocol*."],
    ["Borrowed from process control, 1970s.", "Introduced by Anthropic, 2024."],
    ["Answers: *how do I decide what to do next?*",
     "Answers: *how does my agent reach a tool that lives somewhere else?*"],
    ["Changes your agent’s reasoning.", "Changes your agent’s plumbing."],
], y=2.05, col_w=[1, 1.1], size=14, row_h=0.46)
d.label(s, ML, 4.85, CW,
        "They are orthogonal. You can — and often should — run an "
        "!!MPC-style agent whose tools are served over MCP!!.",
        size=16, color=INK, align=PP_ALIGN.CENTER)

s = d.slide("MCP — the problem it solves")
d.body(s, [
    "So far, every tool has been a Python function in the same process as the loop. That does not "
    "survive contact with a real simulation stack.",
    "",
    "**Your solver is:** a compiled binary. On a cluster. Behind a scheduler. Written by someone "
    "else, in Fortran, in 1994. Licensed per seat.",
], size=15)
cl = [("ReAct agent", 2.45), ("LangGraph app", 3.25), ("Claude / IDE", 4.05)]
sv = [("mesh server", 2.45), ("solver server", 3.25), ("CAD server", 4.05)]
hub = d.box(s, 5.95, 2.45, 1.55, 2.20, "MCP", FILL_GRAY, size=15)
for nm, yy in cl:
    b = d.box(s, 1.55, yy, 2.35, 0.60, nm, FILL_BLUE, size=11.5, bold_first=False)
    d.connect(s, b, hub, 3, 1)
for nm, yy in sv:
    b = d.box(s, 9.55, yy, 2.35, 0.60, nm, FILL_GREEN, size=11.5, bold_first=False)
    d.connect(s, hub, b, 3, 1)
d.body(s, ["MCP standardizes the contract between the two sides: !!write your solver’s tool "
           "surface once!!, and any MCP-speaking client can drive it — your notebook today, a "
           "colleague’s LangGraph pipeline tomorrow, a desktop assistant after that."],
       y=5.15, size=15)

s = d.slide("MCP — syntax")
d.body(s, ["Server side: decorate the functions you want to expose."], size=15)
d.code(s, '''
from mcp.server.fastmcp import FastMCP

mcp = FastMCP("fem-tools")

@mcp.tool()
def solve_poisson(mesh_id: int) -> dict:
    """Solve -Laplacian(u) = 1 with u=0 on the boundary."""
    u, basis = solve_poisson_on(MESH_REGISTRY[mesh_id])
    return {"dofs": int(basis.N), "H1_error": h1_error(...)}

@mcp.tool()
def refine_by_threshold(mesh_id: int, theta: float = 0.5) -> dict:
    """Doerfler-mark and refine. theta in (0,1)."""
    ...

mcp.run()          # now speaks MCP over stdio or HTTP
''', y=1.80, size=12.5)
d.body(s, ["The type hints and docstrings *become* the JSON schema. Notice this is the same "
           "information as the “what a tool is” slide — MCP does not change what a tool is, only "
           "!!who can reach it!!."], y=5.35, size=15)

s = d.slide("Choosing: a decision guide")
d.table(s, [
    ["If this is true of your problem…", "start here"],
    ["I can’t specify the procedure in advance, and mistakes are cheap", "**ReAct**"],
    ["Something must *always* happen (verify, log, approve), or a human signs off on the result", "**LangGraph**"],
    ["The roles need different tools and different incentives, and I want the disagreement on record", "**AutoGen**"],
    ["Actions are expensive or irreversible, and I own a forward model", "**MPC**"],
    ["My tools live in another process, language, or machine — or others need to reuse them", "**MCP** (with any of the above)"],
], y=1.40, col_w=[2.3, 1.0], size=14, row_h=0.52)
d.label(s, ML, 5.05, CW, "!!Start at the top.!!  Every row down is more machinery, more failure "
                         "surface, and more code you own. Earn each step.",
        size=17, color=INK, align=PP_ALIGN.CENTER)

s = d.slide("A word on framework enthusiasm")
d.body(s, ["All four patterns are ~30 lines of Python. The frameworks add persistence, retries, "
           "streaming, tracing, and a community — real value, at the cost of an abstraction "
           "between you and the loop."], size=15.5)
d.note(s, "**Suggested order of operations.**   1. Write the loop yourself, once, for your "
          "problem. It is an afternoon.   2. Find out which part actually hurts (state? retries? "
          "observability?).   3. *Then* adopt the framework that fixes that specific pain.",
       y=2.60, x=1.2, w=CW - 1.6, size=14.5)
d.body(s, ["In today’s notebook every pattern appears !!twice!!: once in the real framework’s "
           "syntax, and once in ~15 lines of plain Python. The second version is there to "
           "demystify the first — and so the notebook still runs when a `pip install` fails on "
           "conference wifi."], y=4.45, size=15.5)

# ===========================================================  38-44. TOOL SURFACE
divider("Scientific computing\nas a tool surface")

s = d.slide("The workflow you already know")
stg = [("CAD", FILL_GRAY), ("mesh", FILL_BLUE), ("solve", FILL_GREEN), ("post", FILL_ORANGE)]
bw, gap, bx0, by, bh = 2.40, 0.83, ML, 1.32, 0.80
sh = []
for i, (nm, fill) in enumerate(stg):
    sh.append(d.box(s, bx0 + i * (bw + gap), by, bw, bh, nm, fill, size=15))
for i in range(3):
    d.connect(s, sh[i], sh[i + 1], 3, 1)
# feedback loop, routed *below* the row so it never collides with the title
_pc, _mc, _yl = bx0 + 3 * (bw + gap) + bw / 2, bx0 + (bw + gap) + bw / 2, 2.62
for _seg in [d.arrow(s, _pc, by + bh, _pc, _yl, head=False, color=LGRAY),
             d.arrow(s, _pc, _yl, _mc, _yl, head=False, color=LGRAY),
             d.arrow(s, _mc, _yl, _mc, by + bh, color=LGRAY)]:
    _seg.line.dash_style = 4
d.label(s, ML, 2.66, CW, "refine and repeat", size=11)
d.body(s, [
    "Most of you do this daily. The question for the rest of the morning is not how it works — it is:",
], y=3.22, size=15.5)
d.label(s, ML, 3.80, CW, "What does this look like when an LLM is holding the mouse?",
        size=21, color=ACCENT, align=PP_ALIGN.CENTER)
d.body(s, [
    "Three things have to change:",
    "1.  Every stage needs a **JSON-callable signature**.",
    "2.  Objects that can’t cross a JSON boundary (meshes, bases, matrices) need **handles**.",
    "3.  Every stage needs **guardrails**, because the caller is nondeterministic.",
], y=4.60, size=15.5)

s = d.slide("scikit-fem: the whole workflow in 30 lines")
d.code(s, '''
from skfem import *
from skfem.helpers import dot, grad

@BilinearForm                       # weak form:  a(u,v) = int grad u . grad v
def stiffness(u, v, w):
    return dot(grad(u), grad(v))

@LinearForm                         # load:       L(v)   = int 1 * v
def unit_load(v, w):
    return 1.0 * v

mesh  = MeshTri.init_lshaped().refined(2)      # "CAD" + mesh
basis = Basis(mesh, ElementTriP1())            # discretization

K = stiffness.assemble(basis)                  # assemble
b = unit_load.assemble(basis)
u = solve(*condense(K, b, D=mesh.boundary_nodes()))   # solve

mesh.draw(); basis.plot(u)                     # post
''', y=1.30, size=12.5)
d.body(s, ["Pure Python, pip-installable, runs in Colab. Small enough to read in one sitting, "
           "real enough to show genuine convergence behaviour — which is why we use it rather "
           "than a production code."], y=5.75, size=14)

s = d.slide("Re-cutting the workflow as tools")
d.body(s, ["The agent cannot hold a `MeshTri` object. It can hold an integer."], size=15.5)
d.code(s, '''
MESH_REGISTRY = {}                # int -> mesh.  The agent sees only the int.
NEXT_ID = [0]

def _register(mesh):
    mid = NEXT_ID[0]; NEXT_ID[0] += 1
    MESH_REGISTRY[mid] = mesh
    return mid

def tool_solve_poisson(mesh_id: int):
    m = MESH_REGISTRY.get(mesh_id)
    if m is None:
        return {"error": f"no mesh with id {mesh_id}"}     # observation, not crash
    u, basis = solve_poisson_on(m)
    return {"mesh_id": mesh_id, "dofs": int(basis.N),
            "H1_error": h1_error(m, u)}
''', y=1.72, size=12)
d.body(s, ["This !!handle pattern!! is the single most reusable idea in the talk. Meshes, solver "
           "contexts, CAD bodies, open files, database cursors — anything with identity and no "
           "JSON representation gets an integer and a registry."], y=5.40, size=15)

s = d.slide("Task: adaptive refinement on the L-shape")
d.body(s, [
    "Poisson with uniform forcing:",
    "        −Δu = 1  in Ω,        u = 0  on ∂Ω.",
    "",
    "The re-entrant corner gives a singular solution",
    "        u ∼ r^(2/3) · sin(2θ/3),",
    "so u ∈ H^(5/3−ε) but u ∉ H².",
    "",
    "**P1 elements, expected rates:**",
    "- **Uniform** refinement: O(h^(2/3)) asymptotically. Pre-asymptotically you often see 0.7–0.95.",
    "- **Adaptive** refinement: optimal O(N^(−1/2)), i.e. rate 1 against h ∼ N^(−1/2).",
], w=7.7, size=14.5)
d.label(s, ML, 5.15, 7.7, "The gap between those two rates is the entire point of adaptive "
                          "meshing. !!We are going to make an agent rediscover it.!!",
        size=14, color=INK, align=PP_ALIGN.LEFT)
# L-shaped domain, as an editable freeform
x0, y0, U = 8.85, 1.85, 1.42
pts = [(x0 + U, y0 + U), (x0 + 2 * U, y0 + U), (x0 + 2 * U, y0),
       (x0, y0), (x0, y0 + 2 * U), (x0 + U, y0 + 2 * U)]
ff = s.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]))
ff.add_line_segments([(Inches(px), Inches(py)) for px, py in pts[1:]], close=True)
poly = ff.convert_to_shape()
poly.fill.solid(); poly.fill.fore_color.rgb = FILL_BLUE
poly.line.color.rgb = INK; poly.line.width = Pt(1.75)
poly.shadow.inherit = False
dot_ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x0 + U - 0.055), Inches(y0 + U - 0.055),
                          Inches(0.11), Inches(0.11))
dot_.fill.solid(); dot_.fill.fore_color.rgb = ACCENT
dot_.line.fill.background(); dot_.shadow.inherit = False
d.label(s, x0 + U + 0.12, y0 + U + 0.10, 1.6, "re-entrant\ncorner", size=10.5, align=PP_ALIGN.LEFT)
d.label(s, x0 + 0.30, y0 + 0.55, 0.8, "Ω", size=15, color=INK)

s = d.slide("What “adaptive” actually means")
d.body(s, ["**1. A local error indicator ηₖ.**  For each element K, a computable proxy for how "
           "much error lives there. We use the gradient-jump indicator:"], size=15)
s.shapes.add_picture(EQ, Inches(3.9), Inches(2.05), height=Inches(0.78))
d.body(s, ["P1 gradients are discontinuous across element edges, and the size of the jump "
           "correlates with the local error."], y=3.00, size=14.5)
d.body(s, [
    "**2. Dörfler marking.**  Pick a bulk fraction θ ∈ (0,1). Mark the smallest set M whose "
    "squared indicators capture a fraction θ of the total:",
    "",
    "        Σ(K ∈ M)  η²ₖ      ≥      θ · Σ(K)  η²ₖ",
], y=3.65, size=15)
d.note(s, "Sort descending, walk the list until the running sum hits θ · total, refine what you "
          "walked past. θ = 0.5 is the textbook default; smaller θ focuses on the worst offenders, "
          "larger θ approaches uniform refinement.", y=5.45, size=13)

s = d.slide("Tools the FEM agent gets")
d.body(s, ["Each takes a `mesh_id` and returns a dict."], size=15)
d.body(s, [
    "- `solve_poisson(mesh_id)`  →  `{dofs, H1_error}`",
    "- `local_indicators(mesh_id)`  →  summary stats of ηₖ",
    "- `uniform_refine(mesh_id)`  →  new mesh id",
    "- `refine_by_threshold(mesh_id, theta)`  →  Dörfler-mark and refine",
    "- `check_convergence_rate()`  →  log-log slope of the last 3 solves",
    "- `stop(reason)`",
], y=1.70, size=15)
d.body(s, [
    "**Guardrails, in the tools:**",
    "- `MAX_DOFS = 20000` — a refinement that would exceed it returns an error observation instead of allocating.",
    "- Every dispatch is `try/except` wrapped.",
    "- `stop` is the only terminal tool.",
], y=4.15, size=15)

# ===========================================================  45-49. RELIABILITY
divider("Reliability", "or: agents are confidently wrong")

s = d.slide("The failure mode that should actually worry you")
d.body(s, ["A crashed tool call is *fine*. You see a traceback; the model retries."], size=16)
d.label(s, ML, 1.85, CW,
        "The dangerous output is a !!plausible number with a confident summary!! and a bug "
        "upstream of both.", size=19, color=INK, align=PP_ALIGN.CENTER)
d.body(s, [
    "In our setting, all of these produce clean-looking convergence plots:",
    "- the agent refines toward the wrong corner and the error still decreases;",
    "- the error *estimate* is fine but the solver has a sign error in the load;",
    "- the reference energy is computed on too coarse a mesh, so every error looks small;",
    "- the agent stops early and reports the tolerance as met.",
], y=3.05, size=15)
d.label(s, ML, 5.65, CW, "None of these announce themselves. !!Something independent has to check.!!",
        size=17, color=INK, align=PP_ALIGN.CENTER)

s = d.slide("The verifier pattern")
va = d.box(s, 1.55, 1.55, 3.0, 1.05, "Agent A\ndoes the work", FILL_BLUE, size=14)
vb = d.box(s, 5.15, 1.55, 3.0, 1.05, "Agent B\nchecks the claim", FILL_GREEN, size=14)
vc = d.box(s, 8.75, 1.55, 3.0, 1.05, "PASS / FAIL\n+ justification", FILL_ORANGE, size=14)
d.connect(s, va, vb, 3, 1); d.connect(s, vb, vc, 3, 1)
d.label(s, 4.35, 1.80, 1.6, "claim", size=10.5)
d.label(s, 5.15, 2.72, 3.0, "different system prompt\ndifferent tools\ndoes not see A’s reasoning",
        size=10.5, italic=True)
d.body(s, [
    "A verifier is just another agent. What makes it useful is what it is *denied*:",
    "- It does not see how the answer was produced — only the claim.",
    "- It has its !!own tools!!, ideally sharing no code path with A’s.",
    "- Its system prompt rewards finding problems, not agreeing.",
], y=3.85, size=15)
d.note(s, "If A and B share the buggy solver, B will happily confirm the bug. Independence is a "
          "property of the *tools*, not the prompt.", y=5.85, size=13)

s = d.slide("Two verifiers we actually run")
d.body(s, [
    "**Verifier 1 — manufactured solutions (for the FEM task).**",
    "Pick a smooth u_exact, compute f = −Δu_exact symbolically, solve with u_h = u_exact on ∂Ω, "
    "and measure the error. For smooth u_exact and P1 elements the rates are known: "
    "L² error O(h²), H¹-seminorm error O(h).",
    "",
    "The agent *chooses* u_exact — itself a small test of whether it understands the problem. A "
    "verifier that picks u = x + y has proven nothing, since P1 reproduces it exactly.",
    "",
    "**Verifier 2 — an oracle integrator (for the projectile task).**",
    "The simulator uses `odeint` at default tolerance. The verifier re-integrates the claimed "
    "optimum with `solve_ivp(method=\"RK45\", rtol=1e-10)` and reports the discrepancy — catching "
    "the case where the “optimum” is an artifact of integration error rather than physics.",
], size=14.5)
d.label(s, ML, 5.95, CW, "!!Both are ordinary numerical-methods practice.!!  The agent framing "
                         "changes nothing about what makes a result trustworthy.",
        size=15, color=INK, align=PP_ALIGN.CENTER)

s = d.slide("Where verification belongs")
d.body(s, ["Recall the LangGraph exemplar. This is why it was the exemplar."], size=16)
d.table(s, [
    ["Verification as…", "What you get"],
    ["a line in the system prompt", "a suggestion, followed most of the time"],
    ["a tool the agent may call", "used when the agent feels uncertain"],
    ["**a node in a graph**", "**runs every time, or the result never ships**"],
], y=1.95, col_w=[1, 1.35], size=14.5, row_h=0.46)
d.body(s, [
    "The pattern and the reliability requirement are the same decision. Choosing LangGraph over "
    "ReAct for the production FEM workflow *is* choosing to make verification non-optional.",
], y=4.30, size=15.5)
d.note(s, "This is the thread that connects the framework zoo to the engineering: you don’t pick "
          "a framework by taste, you pick it by which guarantee you need.", y=5.45, size=13.5)

# ===========================================================  50-53. HANDS ON
s = d.slide("Hands-on: what you’ll build")
d.body(s, [
    "**Notebook `01_agentic_patterns.ipynb`** — the zoo",
    "- §1–§5: each pattern, twice (framework syntax + ~15 lines from scratch)",
    "- §6: **ReAct vs. MPC head-to-head** on the projectile, same budget",
    "- §7: `scipy.optimize` baseline — the result that keeps us honest",
    "",
    "**Notebook `02_agent_hackathon.ipynb`** — the FEM hackathon",
    "- **A1**: uniform-refinement agent (`uniform_refine` only)",
    "- **A2**: adaptive agent (Dörfler marking, θ = 0.5)",
    "- **A3**: rate-checking agent (adapts θ when the rate drops)",
    "- **V**: manufactured-solution verifier",
], size=15)
d.label(s, ML, 5.55, CW, "Every TODO has a worked solution one cell below. !!Use it.!!  Reading a "
                         "good system prompt is the fastest way to learn to write one.",
        size=15, color=INK, align=PP_ALIGN.CENTER)

s = d.slide("Stretch goals")
d.body(s, ["If you finish early, or want somewhere to go afterward:"], size=15.5)
d.body(s, [
    "- **Swap the backend.** Point the shim at Anthropic or OpenAI and see what changes in the tool-call round trip. (Answer: less than you’d expect.)",
    "- **Break a tool on purpose.** Introduce a sign error in the load vector. Does your verifier catch it? Does the convergence plot?",
    "- **Wire the FEM tools as a real graph.** Take the A2 agent and put the verifier on a conditional edge so a FAIL routes to further refinement.",
    "- **Adversarial verifier.** Rewrite the verifier’s system prompt to reward finding problems. Compare its pass rate against the neutral one.",
    "- **Serve the FEM tools over MCP** and drive them from a client that isn’t this notebook.",
], y=1.85, size=15)

s = d.slide("Taking it further")
d.body(s, [
    "**Agent patterns**",
    "- ReAct — Yao et al. 2022, arxiv.org/abs/2210.03629",
    "- LangGraph — langchain-ai.github.io/langgraph",
    "- AutoGen — microsoft.github.io/autogen",
    "- Model Context Protocol — modelcontextprotocol.io",
    "- smolagents (minimal framework) — huggingface.co/docs/smolagents",
    "",
    "**Agents doing science**",
    "- FunSearch — *Nature* 625, 2024;   AlphaEvolve — DeepMind, 2025",
    "- ChemCrow — arxiv.org/abs/2304.05376",
    "- Coscientist — *Nature* 624, 2023",
    "- Sakana AI Scientist — sakana.ai/ai-scientist",
    "",
    "**Numerics**",
    "- scikit-fem — scikit-fem.readthedocs.io",
    "- Dörfler, *SIAM J. Numer. Anal.* 33(3), 1996",
], size=13.5, spacing=2)

s = d.slide(footer=False)
tb = s.shapes.add_textbox(Inches(ML), Inches(2.6), Inches(CW), Inches(2.0))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
add_runs(p, "Let’s go.", size=44, bold=True)
p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; set_space(p2, before=18)
add_runs(p2, "Open 01_agentic_patterns.ipynb and run Part 0.", size=19)
d.label(s, ML, 5.6, CW, "Materials: github.com/natrask/AESCAPE", size=12)

# ===========================================================  SAVE
out = r"c:\Users\nattr\OneDrive\Desktop\GitRepos\AESCAPE\slides\agentic_ai_intro.pptx"
n = d.save(out)
print(f"wrote {out}")
print(f"{n} slides")
