"""Helpers for building an editable AESCAPE deck with python-pptx."""
import re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

# --- geometry ---------------------------------------------------------------
SW, SH = 13.333, 7.5
ML, MR = 0.62, 0.62
CW = SW - ML - MR
TITLE_TOP, TITLE_H = 0.30, 0.70
BODY_TOP = 1.18
BODY_BOT = 6.82
FOOTER_TOP = 6.92

# --- palette ----------------------------------------------------------------
INK    = RGBColor(0x1A, 0x1A, 0x1A)
GRAY   = RGBColor(0x6E, 0x6E, 0x6E)
LGRAY  = RGBColor(0x9A, 0x9A, 0x9A)
ACCENT = RGBColor(0xAA, 0x1E, 0x1E)
KW     = RGBColor(0x00, 0x5A, 0xA0)
STR    = RGBColor(0x8C, 0x3C, 0x14)
CMT    = RGBColor(0x78, 0x78, 0x78)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

FILL_BLUE   = RGBColor(0xE8, 0xEE, 0xF7)
FILL_ORANGE = RGBColor(0xFA, 0xED, 0xE0)
FILL_GREEN  = RGBColor(0xE6, 0xF2, 0xE6)
FILL_PURPLE = RGBColor(0xF1, 0xE9, 0xF1)
FILL_GRAY   = RGBColor(0xED, 0xED, 0xED)
FILL_CODE   = RGBColor(0xF7, 0xF7, 0xF7)
FILL_NOTE   = RGBColor(0xF2, 0xF4, 0xF7)
EDGE        = RGBColor(0xB0, 0xB0, 0xB0)

FONT   = "Calibri"
MONO   = "Consolas"

PY_KW = set("""def return for in if else elif while try except class import from as with
yield lambda None True None and or not pass break continue await async global assert""".split())


# --- low-level xml helpers ---------------------------------------------------
def _el(tag, **attrs):
    e = etree.SubElement(etree.Element("tmp"), qn(tag))
    for k, v in attrs.items():
        e.set(k, v)
    return e


def set_bullet(p, char="▪", color=None, size_pct=90):
    """Give a paragraph a real PowerPoint bullet (so it stays editable)."""
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for old in pPr.findall(qn(tag)):
            pPr.remove(old)
    font = etree.SubElement(pPr, qn("a:buFont"))
    font.set("typeface", "Arial")
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", char)
    if color is not None:
        clr = etree.SubElement(pPr, qn("a:buClr"))
        srgb = etree.SubElement(clr, qn("a:srgbClr"))
        srgb.set("val", str(color))
        pPr.insert(0, clr)
    sz = etree.SubElement(pPr, qn("a:buSzPct"))
    sz.set("val", str(size_pct * 1000))


def set_indent(p, left_in, hang_in):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(left_in))))
    pPr.set("indent", str(-Emu(Inches(hang_in))))


def set_space(p, before=0, after=0, line=None):
    if before: p.space_before = Pt(before)
    if after:  p.space_after = Pt(after)
    if line:   p.line_spacing = line


def add_slidenum_field(tf_para, size, color):
    """A live slide-number field: survives reordering."""
    fld = etree.SubElement(tf_para._p, qn("a:fld"))
    fld.set("id", "{B7C7A0D1-1A2B-4C3D-9E5F-6A7B8C9D0E1F}")
    fld.set("type", "slidenum")
    rPr = etree.SubElement(fld, qn("a:rPr"))
    rPr.set("lang", "en-US")
    rPr.set("sz", str(int(size * 100)))
    fill = etree.SubElement(rPr, qn("a:solidFill"))
    srgb = etree.SubElement(fill, qn("a:srgbClr"))
    srgb.set("val", str(color))
    latin = etree.SubElement(rPr, qn("a:latin"))
    latin.set("typeface", FONT)
    t = etree.SubElement(fld, qn("a:t"))
    t.text = "1"


def set_arrow(conn, head=True, tail=False, width=1.25, color=GRAY):
    ln = conn.line._get_or_add_ln()
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    if head:
        e = etree.SubElement(ln, qn("a:tailEnd"))
        e.set("type", "triangle"); e.set("w", "med"); e.set("len", "med")
    if tail:
        e = etree.SubElement(ln, qn("a:headEnd"))
        e.set("type", "triangle"); e.set("w", "med"); e.set("len", "med")


# --- inline markup -----------------------------------------------------------
TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`|!!.+?!!|\*[^*]+?\*)")

def add_runs(p, text, size=16, color=INK, font=FONT, bold=False):
    """Inline markup:  **bold**  `code`  *italic*  !!accent!!"""
    for piece in TOKEN.split(text):
        if not piece:
            continue
        r = p.add_run()
        r.font.size = Pt(size)
        r.font.name = font
        r.font.color.rgb = color
        r.font.bold = bold
        if piece.startswith("**") and piece.endswith("**"):
            r.text = piece[2:-2]; r.font.bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            r.text = piece[1:-1]; r.font.name = MONO; r.font.size = Pt(size - 1.5)
        elif piece.startswith("!!") and piece.endswith("!!"):
            r.text = piece[2:-2]; r.font.bold = True; r.font.color.rgb = ACCENT
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            r.text = piece[1:-1]; r.font.italic = True
        else:
            r.text = piece
    return p


# --- syntax colouring --------------------------------------------------------
CODE_TOK = re.compile(r"(#.*$)|(\"[^\"]*\"|'[^']*')|(\b[A-Za-z_][A-Za-z_0-9]*\b)")

def add_code_runs(p, line, size):
    pos = 0
    def emit(txt, color, italic=False):
        if not txt:
            return
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size); r.font.name = MONO
        r.font.color.rgb = color; r.font.italic = italic
    for m in CODE_TOK.finditer(line):
        emit(line[pos:m.start()], INK)
        cmt, s, word = m.group(1), m.group(2), m.group(3)
        if cmt is not None:
            emit(cmt, CMT, italic=True)
        elif s is not None:
            emit(s, STR)
        elif word in PY_KW:
            emit(word, KW)
        else:
            emit(word, INK)
        pos = m.end()
    emit(line[pos:], INK)


class Deck:
    def __init__(self, footer="AESCAPE 2026  ·  Agentic AI"):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SW)
        self.prs.slide_height = Inches(SH)
        self.footer = footer
        self.blank = self.prs.slide_layouts[6]

    # -- slide scaffolding ---------------------------------------------------
    def slide(self, title=None, footer=True, notes=None):
        s = self.prs.slides.add_slide(self.blank)
        if title is not None:
            self.title(s, title)
        if footer:
            self.add_footer(s)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def title(self, s, text, size=25):
        tb = s.shapes.add_textbox(Inches(ML), Inches(TITLE_TOP),
                                  Inches(CW), Inches(TITLE_H))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        add_runs(p, text, size=size, bold=True)
        return tb

    def add_footer(self, s):
        tb = s.shapes.add_textbox(Inches(SW - MR - 4.2), Inches(FOOTER_TOP),
                                  Inches(4.2), Inches(0.32))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run(); r.text = self.footer + "  ·  "
        r.font.size = Pt(9); r.font.color.rgb = LGRAY; r.font.name = FONT
        add_slidenum_field(p, 9, "9A9A9A")
        return tb

    # -- content blocks ------------------------------------------------------
    def body(self, s, lines, x=ML, y=BODY_TOP, w=None, size=16, spacing=9,
             line_spacing=1.0):
        """lines: list of strings. '- ' = bullet, '  - ' = sub-bullet, '' = gap."""
        w = w or CW
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        first = True
        for raw in lines:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            if raw.strip() == "":
                p.text = ""
                r = p.add_run(); r.text = " "; r.font.size = Pt(size // 2)
                continue
            if raw.startswith("  - "):
                set_bullet(p, "–", color="6E6E6E")
                set_indent(p, 0.66, 0.24)
                add_runs(p, raw[4:], size=size - 1.5)
            elif raw.startswith("- "):
                set_bullet(p, "▪", color="AA1E1E")
                set_indent(p, 0.30, 0.30)
                add_runs(p, raw[2:], size=size)
            elif raw.startswith("# "):
                add_runs(p, raw[2:], size=size + 1, bold=True)
            else:
                add_runs(p, raw, size=size)
            set_space(p, after=spacing)
            p.line_spacing = line_spacing
        return tb

    def code(self, s, code_text, x=ML, y=BODY_TOP, w=None, size=11.5, pad=0.14):
        w = w or CW
        lines = code_text.strip("\n").split("\n")
        h = pad * 2 + len(lines) * (size * 1.30 / 72.0)
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                 Inches(w), Inches(h))
        box.adjustments[0] = 0.03
        box.fill.solid(); box.fill.fore_color.rgb = FILL_CODE
        box.line.color.rgb = RGBColor(0xE2, 0xE2, 0xE2); box.line.width = Pt(0.75)
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.14); tf.margin_right = Inches(0.08)
        tf.margin_top = Inches(pad); tf.margin_bottom = Inches(pad)
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT      # autoshapes default to centred
            p.line_spacing = 1.06
            set_space(p, after=0)
            add_code_runs(p, line, size)
        return box

    def note(self, s, text, x=ML, y=None, w=None, size=13, fill=True):
        """The small grey 'block' callouts from the beamer deck."""
        w = w or CW
        y = y if y is not None else BODY_BOT - 0.9
        h = 0.34 + 0.235 * (1 + len(text) // int(w * 15))
        if fill:
            box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                     Inches(w), Inches(h))
            box.adjustments[0] = 0.06
            box.fill.solid(); box.fill.fore_color.rgb = FILL_NOTE
            box.line.color.rgb = RGBColor(0xD6, 0xDD, 0xE8); box.line.width = Pt(0.75)
            box.shadow.inherit = False
            tf = box.text_frame
        else:
            tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.16)
        tf.margin_top = tf.margin_bottom = Inches(0.09)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.paragraphs[0].alignment = PP_ALIGN.LEFT
        add_runs(tf.paragraphs[0], text, size=size)
        return tf

    def table(self, s, rows, x=ML, y=BODY_TOP, w=None, col_w=None, size=13,
              header=True, row_h=0.34):
        w = w or CW
        nr, nc = len(rows), len(rows[0])
        shape = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w),
                                   Inches(row_h * nr))
        tbl = shape.table
        tbl.first_row = header
        tbl.horz_banding = False
        if col_w:
            total = sum(col_w)
            for i, cwid in enumerate(col_w):
                tbl.columns[i].width = Emu(int(Inches(w) * cwid / total))
        for ri, row in enumerate(rows):
            tbl.rows[ri].height = Inches(row_h)
            for ci, cell_text in enumerate(row):
                cell = tbl.cell(ri, ci)
                cell.margin_left = Inches(0.09); cell.margin_right = Inches(0.06)
                cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                cell.fill.fore_color.rgb = (RGBColor(0xEC, 0xEF, 0xF3) if (header and ri == 0)
                                            else WHITE)
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT
                add_runs(p, cell_text, size=size, bold=(header and ri == 0))
        return shape

    # -- diagram primitives --------------------------------------------------
    def box(self, s, x, y, w, h, text, fill=FILL_BLUE, size=11, bold_first=True,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE):
        sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        try:
            sh.adjustments[0] = 0.10
        except Exception:
            pass
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
        sh.line.color.rgb = EDGE; sh.line.width = Pt(1.0)
        sh.shadow.inherit = False
        tf = sh.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            set_space(p, after=0)
            add_runs(p, line, size=(size if i == 0 else size - 1.5),
                     bold=(bold_first and i == 0),
                     color=(INK if i == 0 else GRAY))
        return sh

    def connect(self, s, a, b, a_pt=3, b_pt=1, kind=MSO_CONNECTOR.STRAIGHT,
                color=GRAY, width=1.25):
        """Glued arrow: stays attached when the boxes are moved."""
        conn = s.shapes.add_connector(kind, Inches(1), Inches(1), Inches(2), Inches(2))
        conn.begin_connect(a, a_pt)
        conn.end_connect(b, b_pt)
        set_arrow(conn, color=color, width=width)
        return conn

    def arrow(self, s, x1, y1, x2, y2, color=GRAY, width=1.25,
              kind=MSO_CONNECTOR.STRAIGHT, head=True, tail=False):
        conn = s.shapes.add_connector(kind, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        set_arrow(conn, head=head, tail=tail, color=color, width=width)
        return conn

    def label(self, s, x, y, w, text, size=10, color=GRAY, align=PP_ALIGN.CENTER,
              italic=False):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            set_space(p, after=0)
            add_runs(p, line, size=size, color=color)
            if italic:
                for r in p.runs:
                    r.font.italic = True
        return tb

    def save(self, path):
        self.prs.save(path)
        return len(self.prs.slides._sldIdLst)
